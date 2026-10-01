"""Ready-made document types to start from, for a company that has no template of its own yet:
a site safety inspection, a work order, a site visit report and an incident report. Each is generated as a Word file with
the company's logo and name (through {{ branding }}), so it is theirs from the first report, and
in the company's language. A company with its own document uses the template assistant instead."""

from collections.abc import Callable
from dataclasses import dataclass
from typing import Any

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor
from docx.table import _Cell

from app.services.i18n import normalize_language

_INK = RGBColor(0x1F, 0x2A, 0x44)
_MUTED = RGBColor(0x5B, 0x63, 0x70)

Text = dict[str, str]  # one per language


@dataclass(frozen=True)
class StarterField:
    name: str
    type: str
    label: Text
    description: Text
    required: bool = True
    options: dict[str, list[str]] | None = None
    columns: dict[str, tuple[str, Text]] | None = None  # column → (type, description)
    multiple: bool | None = None

    def spec(self, language: str) -> dict[str, Any]:
        spec: dict[str, Any] = {
            "type": self.type,
            "label": self.label[language],
            "description": self.description[language],
            "required": self.required,
        }
        if self.options:
            spec["options"] = self.options[language]
        if self.columns:
            spec["columns"] = {name: {"type": kind, "description": text[language]} for name, (kind, text) in self.columns.items()}
        if self.multiple is not None:
            spec["multiple"] = self.multiple
        return spec


@dataclass(frozen=True)
class Starter:
    key: str
    name: Text
    description: Text
    instructions: Text
    fields: tuple[StarterField, ...]
    layout: Callable[["_Writer", str], None]

    def field_schema(self, language: str) -> dict[str, Any]:
        return {f.name: f.spec(language) for f in self.fields}


class _Writer:
    """Small helpers for a clean, neutral document the company can restyle in Word."""

    def __init__(self, title: str, a4: bool) -> None:
        self.document: DocxDocument = Document()
        styles = self.document.styles
        styles["Normal"].font.name = "Calibri"
        styles["Normal"].font.size = Pt(10.5)
        section = self.document.sections[0]
        if a4:  # python-docx starts on US Letter; Spain prints on A4
            section.page_width, section.page_height = Mm(210), Mm(297)
        section.top_margin = section.bottom_margin = Inches(0.7)
        section.left_margin = section.right_margin = Inches(0.8)
        header = section.header.paragraphs[0]
        header.add_run("{{ branding.logo }}")
        name = section.header.add_paragraph()
        run = name.add_run("{{ branding.name }}")
        run.font.size, run.font.color.rgb = Pt(9), _MUTED
        heading = self.document.add_paragraph()
        heading_run = heading.add_run(title.upper())
        heading_run.bold, heading_run.font.size, heading_run.font.color.rgb = True, Pt(17), _INK
        heading.paragraph_format.space_after = Pt(10)

    def only_if(self, field_name: str) -> None:
        """What follows, until done(), is left out when the report has nothing for it: no empty
        "Photos" heading on a report without photos."""
        self.document.add_paragraph("{%p if " + field_name + " %}")

    def done(self) -> None:
        self.document.add_paragraph("{%p endif %}")

    def section(self, text: str) -> None:
        paragraph = self.document.add_paragraph()
        run = paragraph.add_run(text)
        run.bold, run.font.size, run.font.color.rgb = True, Pt(12), _INK
        paragraph.paragraph_format.space_before, paragraph.paragraph_format.space_after = Pt(10), Pt(3)

    def line(self, *pairs: tuple[str, str]) -> None:
        """Label: value pairs on one line."""
        paragraph = self.document.add_paragraph()
        for i, (label, field_name) in enumerate(pairs):
            paragraph.add_run(("    " if i else "") + f"{label}: ").bold = True
            paragraph.add_run("{{ %s }}" % field_name)  # noqa: UP031 — Jinja braces read better than an f-string's

    def text(self, field_name: str) -> None:
        self.document.add_paragraph("{{ %s }}" % field_name)  # noqa: UP031

    def bullets(self, field_name: str) -> None:
        self.document.add_paragraph("{%p for item in " + field_name + " %}")
        self.document.add_paragraph("{{ item }}", style="List Bullet")
        self.document.add_paragraph("{%p endfor %}")

    def table(self, field_name: str, columns: list[tuple[str, str]]) -> None:
        """A header row, then one row per item (docxtpl's row loop between two marker rows)."""
        table = self.document.add_table(rows=4, cols=len(columns))
        table.style = "Table Grid"
        for cell, (label, _) in zip(table.rows[0].cells, columns, strict=True):
            _shade(cell, "1F2A44")
            run = cell.paragraphs[0].add_run(label)
            run.bold, run.font.color.rgb = True, RGBColor(0xFF, 0xFF, 0xFF)
        table.rows[1].cells[0].paragraphs[0].add_run("{%tr for item in " + field_name + " %}")
        for cell, (_, column) in zip(table.rows[2].cells, columns, strict=True):
            cell.paragraphs[0].add_run("{{ item.%s }}" % column)  # noqa: UP031
        table.rows[3].cells[0].paragraphs[0].add_run("{%tr endfor %}")

    def photos(self, field_name: str) -> None:
        self.document.add_paragraph("{%p for photo in " + field_name + " %}")
        self.document.add_paragraph("{{ photo }}").alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.document.add_paragraph("{%p endfor %}")

    def signatures(self, left: str, right: str) -> None:
        self.document.add_paragraph().paragraph_format.space_before = Pt(28)
        table = self.document.add_table(rows=1, cols=2)
        for cell, label in zip(table.rows[0].cells, (left, right), strict=True):
            cell.paragraphs[0].add_run("_" * 30 + "\n" + label).font.color.rgb = _MUTED


def _shade(cell: _Cell, color: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shading)


def _f(name: str, kind: str, es: tuple[str, str], en: tuple[str, str], **extra: Any) -> StarterField:
    return StarterField(name, kind, {"es": es[0], "en": en[0]}, {"es": es[1], "en": en[1]}, **extra)


_PHOTOS = _f("fotos", "image", ("Fotos", "Fotos enviadas con el informe"), ("Photos", "Photos sent with the report"), required=False, multiple=True)

_WORK_ORDER_FIELDS = (
    _f("cliente", "str", ("Cliente", "Cliente o comunidad atendida"), ("Client", "Client or building served")),
    _f("direccion", "str", ("Dirección", "Dirección donde se hizo el trabajo"), ("Address", "Where the work was done")),
    _f("fecha", "date", ("Fecha", "Día en que se hizo el trabajo"), ("Date", "Day the work was done")),
    _f("hora_llegada", "time", ("Hora de llegada", "Hora de llegada (HH:MM)"), ("Arrival", "Arrival time (HH:MM)"), required=False),
    _f("hora_salida", "time", ("Hora de salida", "Hora de salida (HH:MM)"), ("Departure", "Departure time (HH:MM)"), required=False),
    _f("tecnico", "str", ("Técnico", "Quién hizo el trabajo"), ("Technician", "Who did the work")),
    _f("trabajos", "list[str]", ("Trabajos realizados", "Cada trabajo hecho, uno por elemento"), ("Work done", "Each job done, one per item")),
    _f(
        "materiales", "list[object]", ("Materiales", "Materiales usados"), ("Materials", "Materials used"), required=False,
        columns={
            "material": ("str", {"es": "Qué material", "en": "Which material"}),
            "cantidad": ("float", {"es": "Cuánto", "en": "How much"}),
            "unidad": ("str", {"es": "Unidad (m, ud, kg…)", "en": "Unit (m, pcs, kg…)"}),
        },
    ),
    _f("observaciones", "str", ("Observaciones", "Pendientes, incidencias o avisos al cliente"), ("Notes", "Pending work, issues or notes for the client"), required=False),
    _f("terminado", "bool", ("Trabajo terminado", "Si el trabajo quedó terminado"), ("Finished", "Whether the work was finished")),
    _PHOTOS,
)


def _work_order(w: _Writer, language: str) -> None:
    es = language == "es"
    w.line(("Cliente" if es else "Client", "cliente"))
    w.line(("Dirección" if es else "Address", "direccion"))
    w.line(("Fecha" if es else "Date", "fecha"), ("Llegada" if es else "Arrival", "hora_llegada"), ("Salida" if es else "Departure", "hora_salida"))
    w.line(("Técnico" if es else "Technician", "tecnico"))
    w.section("Trabajos realizados" if es else "Work done")
    w.bullets("trabajos")
    w.only_if("materiales")
    w.section("Materiales" if es else "Materials")
    w.table("materiales", [("Material", "material"), ("Cantidad" if es else "Quantity", "cantidad"), ("Unidad" if es else "Unit", "unidad")])
    w.done()
    w.only_if("observaciones")
    w.section("Observaciones" if es else "Notes")
    w.text("observaciones")
    w.done()
    w.line(("Trabajo terminado" if es else "Finished", "terminado"))
    # Signed under the work, photos after: a photo would otherwise push the signatures alone onto page 2.
    w.signatures("Firma del técnico" if es else "Technician's signature", "Firma del cliente" if es else "Client's signature")
    w.only_if("fotos")
    w.section("Fotos" if es else "Photos")
    w.photos("fotos")
    w.done()


_VISIT_FIELDS = (
    _f("cliente", "str", ("Cliente", "Empresa o persona visitada"), ("Client", "Company or person visited")),
    _f("lugar", "str", ("Lugar", "Dónde fue la visita"), ("Place", "Where the visit was")),
    _f("fecha", "date", ("Fecha", "Día de la visita"), ("Date", "Day of the visit")),
    _f("asistentes", "list[str]", ("Asistentes", "Quién estuvo, con su cargo si se dice"), ("Attendees", "Who was there, with their role if said")),
    _f("motivo", "str", ("Motivo", "Para qué fue la visita"), ("Purpose", "What the visit was for")),
    _f("conclusiones", "str", ("Conclusiones", "Qué se vio, se habló o se acordó"), ("Findings", "What was seen, discussed or agreed")),
    _f("proximos_pasos", "list[str]", ("Próximos pasos", "Lo que queda por hacer, uno por elemento"), ("Next steps", "What is left to do, one per item"), required=False),
    _PHOTOS,
)


def _visit(w: _Writer, language: str) -> None:
    es = language == "es"
    w.line(("Cliente" if es else "Client", "cliente"))
    w.line(("Lugar" if es else "Place", "lugar"), ("Fecha" if es else "Date", "fecha"))
    w.line(("Motivo" if es else "Purpose", "motivo"))
    w.section("Asistentes" if es else "Attendees")
    w.bullets("asistentes")
    w.section("Conclusiones" if es else "Findings")
    w.text("conclusiones")
    w.only_if("proximos_pasos")
    w.section("Próximos pasos" if es else "Next steps")
    w.bullets("proximos_pasos")
    w.done()
    w.only_if("fotos")
    w.section("Fotos" if es else "Photos")
    w.photos("fotos")
    w.done()


_SEVERITY = {"es": ["Baja", "Media", "Alta"], "en": ["Low", "Medium", "High"]}

_INCIDENT_FIELDS = (
    _f("fecha", "date", ("Fecha", "Día de la incidencia"), ("Date", "Day of the incident")),
    _f("hora", "time", ("Hora", "Hora aproximada (HH:MM)"), ("Time", "Approximate time (HH:MM)"), required=False),
    _f("lugar", "str", ("Lugar", "Dónde ocurrió"), ("Place", "Where it happened")),
    _f("comunicado_por", "str", ("Comunicado por", "Quién informa"), ("Reported by", "Who reports it")),
    _f("descripcion", "str", ("Descripción", "Qué ha pasado"), ("Description", "What happened")),
    _f("gravedad", "enum", ("Gravedad", "Baja, media o alta"), ("Severity", "Low, medium or high"), options=_SEVERITY),
    _f("acciones", "list[str]", ("Acciones tomadas", "Qué se hizo ya, una por elemento"), ("Actions taken", "What was done already, one per item"), required=False),
    _f("pendiente", "str", ("Pendiente", "Qué queda por resolver"), ("Pending", "What is left to resolve"), required=False),
    _PHOTOS,
)


def _incident(w: _Writer, language: str) -> None:
    es = language == "es"
    w.line(("Fecha" if es else "Date", "fecha"), ("Hora" if es else "Time", "hora"))
    w.line(("Lugar" if es else "Place", "lugar"))
    w.line(("Comunicado por" if es else "Reported by", "comunicado_por"), ("Gravedad" if es else "Severity", "gravedad"))
    w.section("Descripción" if es else "Description")
    w.text("descripcion")
    w.only_if("acciones")
    w.section("Acciones tomadas" if es else "Actions taken")
    w.bullets("acciones")
    w.done()
    w.only_if("pendiente")
    w.section("Pendiente" if es else "Pending")
    w.text("pendiente")
    w.done()
    w.only_if("fotos")
    w.section("Fotos" if es else "Photos")
    w.photos("fotos")
    w.done()


_SAFETY_VISIT_FIELDS = (
    _f("obra", "str", ("Obra", "Nombre y dirección de la obra"), ("Site", "Site name and address")),
    _f("promotor", "str", ("Promotor", "Promotor de la obra"), ("Client", "Client or developer of the project"), required=False),
    _f("fecha", "date", ("Fecha", "Día de la visita"), ("Date", "Day of the visit")),
    _f("coordinador", "str", ("Coordinador/a", "Quién hace la visita"), ("Inspector", "Who made the visit")),
    _f("fase", "str", ("Fase de obra", "En qué fase está la obra"), ("Stage", "Which stage the works are at"), required=False),
    _f(
        "empresas_presentes", "list[str]",
        ("Empresas presentes", "Contratas y subcontratas trabajando, con su actividad si se dice"),
        ("Contractors on site", "Contractors and subcontractors working, with their trade if said"),
    ),
    _f(
        "deficiencias", "list[object]",
        ("Deficiencias", "Cada deficiencia detectada, una por fila"),
        ("Issues", "Each issue found, one per row"),
        required=False,
        columns={
            "deficiencia": ("str", {"es": "Qué se ha visto", "en": "What was found"}),
            "empresa": ("str", {"es": "Empresa responsable", "en": "Contractor responsible"}),
            "medida": ("str", {"es": "Medida correctora", "en": "Corrective action"}),
            "plazo": ("str", {"es": "Plazo: 'Inmediato', 'Corregido en el momento' o una fecha escrita como 09/10/2026 (no 2026-10-09)",
                              "en": "Deadline: 'Immediate', 'Fixed on the spot' or a date written like Oct 9, 2026 (not 2026-10-09)"}),
        },
    ),
    _f("paralizacion", "bool", ("Paralización de trabajos", "Si se han paralizado trabajos"), ("Work stopped", "Whether any work was stopped"), required=False),
    _f(
        "libro_incidencias", "bool",
        ("Anotado en el libro de incidencias", "Si se ha anotado en el libro de incidencias"),
        ("Recorded in the site log", "Whether it was recorded in the site safety log"),
        required=False,
    ),
    _f("observaciones", "str", ("Observaciones", "Lo demás que se ha visto, bien o mal"), ("Notes", "Anything else seen, good or bad"), required=False),
    _f("proxima_visita", "date", ("Próxima visita", "Fecha de la próxima visita"), ("Next visit", "Date of the next visit"), required=False),
    _PHOTOS,
)


def _safety_visit(w: _Writer, language: str) -> None:
    es = language == "es"
    w.line(("Obra" if es else "Site", "obra"))
    w.line(("Promotor" if es else "Client", "promotor"))
    w.line(("Fecha" if es else "Date", "fecha"), ("Coordinador/a" if es else "Inspector", "coordinador"))
    w.line(("Fase de obra" if es else "Stage", "fase"))
    w.section("Empresas presentes" if es else "Contractors on site")
    w.bullets("empresas_presentes")
    w.only_if("deficiencias")
    w.section("Deficiencias detectadas" if es else "Issues found")
    w.table("deficiencias", [
        ("Deficiencia" if es else "Issue", "deficiencia"),
        ("Empresa" if es else "Contractor", "empresa"),
        ("Medida correctora" if es else "Corrective action", "medida"),
        ("Plazo" if es else "Deadline", "plazo"),
    ])
    w.done()
    w.section("Actuaciones" if es else "Actions")
    w.line(("Paralización de trabajos" if es else "Work stopped", "paralizacion"))
    w.line(("Anotado en el libro de incidencias" if es else "Recorded in the site log", "libro_incidencias"))
    w.only_if("observaciones")
    w.section("Observaciones" if es else "Notes")
    w.text("observaciones")
    w.done()
    w.only_if("proxima_visita")
    w.line(("Próxima visita" if es else "Next visit", "proxima_visita"))
    w.done()
    w.signatures("Coordinador/a de seguridad y salud" if es else "Inspector", "Recibí: jefe/a de obra" if es else "Received: site manager")
    w.only_if("fotos")
    w.section("Fotos" if es else "Photos")
    w.photos("fotos")
    w.done()


STARTERS: dict[str, Starter] = {
    s.key: s
    for s in (
        Starter(
            "safety_visit",
            {"es": "Visita de seguridad y salud", "en": "Site safety inspection"},
            {"es": "Una visita de coordinación a una obra: empresas presentes, deficiencias con su responsable y plazo, paralizaciones.",
             "en": "A safety inspection of a site: contractors on site, issues with who fixes them and by when, work stopped."},
            {"es": "Un coordinador de seguridad y salud dicta la visita al salir de la obra. Cada deficiencia es una fila con "
                   "su empresa responsable, la medida correctora y el plazo ('Inmediato' si debe corregirse ya, 'Corregido en "
                   "el momento' si ya se corrigió). Los nombres de obras, empresas y personas van con sus mayúsculas aunque "
                   "la transcripción los traiga en minúscula.",
             "en": "A safety inspector dictates the visit when leaving the site. Each issue is a row with the contractor "
                   "responsible, the corrective action and the deadline ('Immediate' if it must be fixed now, 'Fixed on the "
                   "spot' if it already was). Names of sites, companies and people keep their capitals even when the "
                   "transcript has them in lowercase."},
            _SAFETY_VISIT_FIELDS,
            _safety_visit,
        ),
        Starter(
            "work_order",
            {"es": "Parte de trabajo", "en": "Work order"},
            {"es": "Lo que hizo un técnico en una visita: trabajos, materiales, horas y fotos.",
             "en": "What a technician did on a call: jobs, materials, hours and photos."},
            {"es": "Un técnico dicta el parte al terminar. 'Hoy' y 'ayer' se refieren a la fecha del mensaje.",
             "en": "A technician dictates the work order when finishing. 'Today' and 'yesterday' are relative to the message date."},
            _WORK_ORDER_FIELDS,
            _work_order,
        ),
        Starter(
            "site_visit",
            {"es": "Informe de visita", "en": "Site visit report"},
            {"es": "Una visita a un cliente u obra: quién estuvo, qué se vio y qué queda pendiente.",
             "en": "A visit to a client or site: who was there, what was seen and what is next."},
            {"es": "Informe de una visita comercial o técnica.", "en": "Report of a sales or technical visit."},
            _VISIT_FIELDS,
            _visit,
        ),
        Starter(
            "incident",
            {"es": "Informe de incidencia", "en": "Incident report"},
            {"es": "Una avería, accidente o queja: qué pasó, cómo de grave y qué se ha hecho.",
             "en": "A breakdown, accident or complaint: what happened, how serious and what was done."},
            {"es": "Informe de una incidencia comunicada por voz.", "en": "Report of an incident told by voice."},
            _INCIDENT_FIELDS,
            _incident,
        ),
    )
}


def describe(language: str) -> list[dict[str, Any]]:
    language = normalize_language(language)
    return [
        {
            "key": s.key,
            "name": s.name[language],
            "description": s.description[language],
            "fields": [f.label[language] for f in s.fields],
        }
        for s in STARTERS.values()
    ]


def build(key: str, language: str, path: str) -> dict[str, Any]:
    """Writes the starter's template to `path`; returns its field schema."""
    language = normalize_language(language)
    starter = STARTERS[key]
    writer = _Writer(starter.name[language], a4=language == "es")
    starter.layout(writer, language)
    writer.document.save(path)
    return starter.field_schema(language)

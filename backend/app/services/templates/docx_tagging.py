"""Turning a filled-in example document into a template, by editing the Word file itself: each
value that changes from one report to the next is replaced by its tag ({{ field }}), repeated
table rows become one row inside a row loop, and a bulleted list becomes one item inside a
paragraph loop. Everything else — fonts, colours, tables, headers, the logo — stays the client's.

Paragraphs are addressed by ids that come from reading the file the same way every time:
b3 (fourth paragraph of the body), t0r2c1p0 (first table, row 2, cell 1, paragraph 0),
h0p1 / f0p1 (header / footer of section 0)."""

import copy
from collections.abc import Iterator
from dataclasses import dataclass, field

from docx import Document
from docx.document import Document as DocxDocument
from docx.table import Table, _Cell, _Row
from docx.text.hyperlink import Hyperlink
from docx.text.paragraph import Paragraph
from docx.text.run import Run

ITEM = "item"  # the loop variable: {{ item }} in a list, {{ item.column }} in a table row


@dataclass(frozen=True)
class Block:
    id: str
    text: str


@dataclass(frozen=True)
class TableShape:
    id: str
    rows: list[list[str]]  # the text of each cell, row by row


@dataclass
class TagPlan:
    """What to tag, as decided in the panel (each name is a field of the document type)."""

    values: list[tuple[str, str, str]] = field(default_factory=list)  # (block id, exact text, field)
    tables: list[tuple[str, int, int, dict[int, str], str]] = field(default_factory=list)  # (table id, first row, last row, cell → column, field)
    lists: list[tuple[list[str], str]] = field(default_factory=list)  # (block ids in order, field)


def _tag(expression: str) -> str:
    return "{{ " + expression + " }}"


def _runs(paragraph: Paragraph) -> list[Run]:
    """Every run, including those inside hyperlinks (paragraph.runs skips them)."""
    runs: list[Run] = []
    for item in paragraph.iter_inner_content():
        runs.extend(item.runs if isinstance(item, Hyperlink) else [item])
    return runs


def paragraph_text(paragraph: Paragraph) -> str:
    return "".join(run.text for run in _runs(paragraph))


def _cells(row: _Row, table: Table) -> list[_Cell]:
    # The physical cells: row.cells repeats a merged cell once per column it spans.
    return [_Cell(tc, table) for tc in row._tr.tc_lst]


def _tables(document: DocxDocument) -> list[Table]:
    return list(document.tables)


def iter_paragraphs(document: DocxDocument) -> Iterator[tuple[str, Paragraph]]:
    for i, paragraph in enumerate(document.paragraphs):
        yield f"b{i}", paragraph
    for t, table in enumerate(_tables(document)):
        for r, row in enumerate(table.rows):
            for c, cell in enumerate(_cells(row, table)):
                for p, paragraph in enumerate(cell.paragraphs):
                    yield f"t{t}r{r}c{c}p{p}", paragraph
    for s, section in enumerate(document.sections):
        for i, paragraph in enumerate(section.header.paragraphs):
            yield f"h{s}p{i}", paragraph
        for i, paragraph in enumerate(section.footer.paragraphs):
            yield f"f{s}p{i}", paragraph


def read_blocks(path: str) -> tuple[list[Block], list[TableShape]]:
    """The text of every non-empty paragraph, and the shape of every table."""
    document = Document(path)
    blocks = [Block(id_, text) for id_, p in iter_paragraphs(document) if (text := paragraph_text(p)).strip()]
    tables = [
        TableShape(
            f"t{t}",
            [[" ".join(paragraph_text(p) for p in cell.paragraphs).strip() for cell in _cells(row, table)] for row in table.rows],
        )
        for t, table in enumerate(_tables(document))
    ]
    return blocks, tables


def replace_text(paragraph: Paragraph, old: str, new: str) -> bool:
    """Replaces the first `old` in the paragraph with `new`, even when Word split it across runs:
    the replacement takes the formatting of the run where `old` starts. False when not found."""
    runs = _runs(paragraph)
    full = "".join(run.text for run in runs)
    start = full.find(old)
    if start < 0 or not old:
        return False
    end = start + len(old)
    position, written = 0, False
    for run in runs:
        text = run.text
        run_start, run_end = position, position + len(text)
        position = run_end
        if run_end <= start or run_start >= end:
            continue
        cut_from, cut_to = max(start, run_start) - run_start, min(end, run_end) - run_start
        run.text = text[:cut_from] + ("" if written else new) + text[cut_to:]
        written = True
    return True


def _set_text(paragraph: Paragraph, text: str) -> None:
    """The whole paragraph becomes `text`, in the formatting of its first run."""
    runs = _runs(paragraph)
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def _remove(element: object) -> None:
    parent = element.getparent()  # type: ignore[attr-defined]
    parent.remove(element)


def _tag_table(table: Table, first: int, last: int, columns: dict[int, str], field_name: str) -> None:
    rows = list(table.rows)
    template_row = rows[first]
    for c, cell in enumerate(_cells(template_row, table)):
        if c in columns:
            _set_text(cell.paragraphs[0], _tag(f"{ITEM}.{columns[c]}"))
            for extra in cell.paragraphs[1:]:
                _remove(extra._p)
    for row in rows[first + 1 : last + 1]:
        _remove(row._tr)
    # docxtpl drops a row whose text is a {%tr %} tag and repeats what lies between the two.
    for tag, place in ((f"{{%tr for {ITEM} in {field_name} %}}", "before"), ("{%tr endfor %}", "after")):
        marker = copy.deepcopy(template_row._tr)
        marker_row = _Row(marker, table)
        for c, cell in enumerate(_cells(marker_row, table)):
            _set_text(cell.paragraphs[0], tag if c == 0 else "")
            for extra in cell.paragraphs[1:]:
                _remove(extra._p)
        if place == "before":
            template_row._tr.addprevious(marker)
        else:
            template_row._tr.addnext(marker)


def _tag_list(paragraphs: list[Paragraph], field_name: str) -> None:
    first = paragraphs[0]
    _set_text(first, _tag(ITEM))
    opening = copy.deepcopy(first._p)
    closing = copy.deepcopy(first._p)
    first._p.addprevious(opening)
    first._p.addnext(closing)
    _set_text(Paragraph(opening, first._parent), f"{{%p for {ITEM} in {field_name} %}}")
    _set_text(Paragraph(closing, first._parent), "{%p endfor %}")
    for paragraph in paragraphs[1:]:
        _remove(paragraph._p)


def apply_plan(source: str, plan: TagPlan, output: str) -> list[str]:
    """Writes the tagged template to `output` and returns what could not be tagged (a value no
    longer found where it was), so the panel can say so instead of silently dropping it."""
    document = Document(source)
    by_id = dict(iter_paragraphs(document))
    tables = _tables(document)
    problems: list[str] = []

    for block_id, text, field_name in plan.values:
        paragraph = by_id.get(block_id)
        if paragraph is None or not replace_text(paragraph, text, _tag(field_name)):
            problems.append(f"«{text}» ({field_name}) was not found where it was")

    # Lists and tables change the structure: done after the values, which address paragraphs by id.
    for block_ids, field_name in plan.lists:
        paragraphs = [by_id[i] for i in block_ids if i in by_id]
        if paragraphs:
            _tag_list(paragraphs, field_name)
        else:
            problems.append(f"The list for {field_name} was not found")

    for table_id, first, last, columns, field_name in plan.tables:
        index = int(table_id.removeprefix("t"))
        if index >= len(tables) or not 0 <= first <= last < len(tables[index].rows):
            problems.append(f"The table for {field_name} was not found")
            continue
        _tag_table(tables[index], first, last, columns, field_name)

    document.save(output)
    return problems


def example_values(source: str, plan: TagPlan) -> dict[str, object]:
    """What the example document said for each field: rendering the template with it should give
    back the example — the clearest proof the tags are in the right places."""
    document = Document(source)
    by_id = dict(iter_paragraphs(document))
    tables = _tables(document)
    values: dict[str, object] = {field_name: text for _, text, field_name in plan.values}
    for block_ids, field_name in plan.lists:
        values[field_name] = [paragraph_text(by_id[i]) for i in block_ids if i in by_id]
    for table_id, first, last, columns, field_name in plan.tables:
        index = int(table_id.removeprefix("t"))
        if index < len(tables):
            table = tables[index]
            values[field_name] = [
                {
                    column: " ".join(paragraph_text(p) for p in cells[c].paragraphs).strip()
                    for c, column in columns.items()
                    if c < len(cells)
                }
                for row in list(table.rows)[first : last + 1]
                for cells in [_cells(row, table)]
            ]
    return values

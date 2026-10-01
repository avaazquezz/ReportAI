"""The template assistant end to end through the API: upload a filled-in example, get a proposal
the model made (cleaned of what cannot be done), preview it, apply it to the document type."""

import io
import uuid
from pathlib import Path
from typing import Any
from unittest.mock import AsyncMock

import pytest
from docx import Document
from docxtpl import DocxTemplate
from httpx import AsyncClient
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.admin import template_assistant as assistant_api
from app.core.config import settings
from app.core.security import hash_password
from app.models.document_template import DocumentTemplate
from app.models.document_type import DocumentType
from app.models.tenant import Tenant
from app.models.tenant_user import TenantUser
from app.services.llm import LLMResult
from app.services.templates import assistant


def _example(tagged: bool = False) -> bytes:
    document = Document()
    document.add_paragraph("PARTE DE TRABAJO")
    if tagged:
        document.add_paragraph("Cliente: {{ cliente }}")
        table = document.add_table(rows=4, cols=2)
        table.rows[0].cells[0].text = "Material"
        table.rows[1].cells[0].text = "{%tr for m in materiales %}"
        table.rows[2].cells[0].text = "{{ m.material }}"
        table.rows[2].cells[1].text = "{{ m.cantidad }}"
        table.rows[3].cells[0].text = "{%tr endfor %}"
    else:
        document.add_paragraph("Cliente: Comunidad Calle Mayor 12")
        document.add_paragraph("Fecha: 14/09/2026")
        table = document.add_table(rows=3, cols=2)
        for r, (a, b) in enumerate([("Material", "Cantidad"), ("Tubo PVC", "3"), ("Silicona", "2")]):
            table.rows[r].cells[0].text, table.rows[r].cells[1].text = a, b
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


MODEL_ANSWER: dict[str, Any] = {
    "fields": [
        {"name": "cliente", "label": "Cliente", "type": "str", "description": "A quién", "required": True},
        {"name": "fecha", "label": "Fecha", "type": "date", "description": "Cuándo", "required": True},
        {"name": "materiales", "label": "Materiales", "type": "list[object]", "description": "Usados",
         "required": False, "columns": [{"name": "material", "type": "str"}, {"name": "cantidad", "type": "int"}]},
        {"name": "Mal Nombre", "label": "x", "type": "str"},  # not a usable name: dropped
        {"name": "fantasma", "label": "Fantasma", "type": "str"},  # nothing points at it: dropped
    ],
    "values": [
        {"block_id": "b1", "text": "Comunidad Calle Mayor 12", "field": "cliente"},
        {"block_id": "b2", "text": "14/09/2026", "field": "fecha"},
        {"block_id": "b2", "text": "15/09/2026", "field": "fecha"},  # not in that paragraph: dropped
        {"block_id": "b0", "text": "PARTE", "field": "fantasma_que_no_existe"},
    ],
    "tables": [{"table_id": "t0", "first_row": 1, "last_row": 2, "field": "materiales",
                "cells": [{"cell": 0, "column": "material"}, {"cell": 1, "column": "cantidad"}]}],
    "lists": [],
}


@pytest.fixture(autouse=True)
def _storage(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(settings, "DOCUMENT_STORAGE_PATH", str(tmp_path))


@pytest.fixture
def model(monkeypatch: pytest.MonkeyPatch) -> AsyncMock:
    ask = AsyncMock(return_value=LLMResult(MODEL_ANSWER, 100, 100, "claude-sonnet-5"))
    monkeypatch.setattr(assistant, "structured_completion", ask)
    return ask


@pytest.fixture
def gotenberg(monkeypatch: pytest.MonkeyPatch) -> AsyncMock:
    async def convert(docx: str, output: str) -> str:
        Path(output).write_bytes(b"%PDF-1.4 " + Path(docx).read_bytes()[:10])
        return output

    mock = AsyncMock(side_effect=convert)
    monkeypatch.setattr(assistant_api, "convert_docx_to_pdf", mock)
    return mock


async def _setup(client: AsyncClient, db: AsyncSession) -> tuple[DocumentType, dict[str, str]]:
    tenant = Tenant(name="Reformas García", slug=f"g-{uuid.uuid4().hex[:6]}", is_active=True)
    db.add(tenant)
    await db.flush()
    user = TenantUser(tenant_id=tenant.id, email="a@garcia.test", hashed_password=hash_password("pw-12345678"),
                      full_name="A", role="tenant_admin", is_active=True)
    doc_type = DocumentType(tenant_id=tenant.id, name="Parte de trabajo", field_schema={"observaciones": {"type": "str"}})
    db.add_all([user, doc_type])
    await db.commit()
    login = await client.post("/auth/login", json={"email": "a@garcia.test", "password": "pw-12345678"})
    return doc_type, {"Authorization": f"Bearer {login.json()['access_token']}"}


async def _upload(client: AsyncClient, doc_type: DocumentType, headers: dict[str, str], content: bytes) -> Any:
    return await client.post(
        f"/document-types/{doc_type.id}/template-assistant",
        files={"file": ("parte.docx", content, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        headers=headers,
    )


async def test_the_proposal_keeps_only_what_can_be_done(client: AsyncClient, db: AsyncSession, model: AsyncMock) -> None:
    doc_type, headers = await _setup(client, db)

    draft = (await _upload(client, doc_type, headers, _example())).json()

    assert "Comunidad Calle Mayor 12" in model.await_args.kwargs["user"]  # the model reads the document
    assert [f["name"] for f in draft["proposal"]["fields"]] == ["cliente", "fecha", "materiales"]
    assert [v["text"] for v in draft["proposal"]["values"]] == ["Comunidad Calle Mayor 12", "14/09/2026"]
    assert draft["blocks"]["b1"] == "Cliente: Comunidad Calle Mayor 12" and draft["already_tagged"] is False


async def test_a_preview_renders_without_saving_anything(
    client: AsyncClient, db: AsyncSession, model: AsyncMock, gotenberg: AsyncMock
) -> None:
    doc_type, headers = await _setup(client, db)
    draft = (await _upload(client, doc_type, headers, _example())).json()

    for mode in ("labels", "example"):
        preview = await client.post(
            f"/document-types/{doc_type.id}/template-assistant/{draft['id']}/preview",
            json={"proposal": draft["proposal"], "mode": mode}, headers=headers,
        )
        assert preview.status_code == 200 and preview.headers["content-type"] == "application/pdf"

    assert (await db.scalars(select(DocumentTemplate))).all() == []


async def test_applying_adds_the_fields_and_activates_the_tagged_template(
    client: AsyncClient, db: AsyncSession, model: AsyncMock
) -> None:
    doc_type, headers = await _setup(client, db)
    draft = (await _upload(client, doc_type, headers, _example())).json()
    proposal = draft["proposal"]
    proposal["fields"][0]["label"] = "Cliente o comunidad"  # corrected in the panel

    applied = await client.post(f"/document-types/{doc_type.id}/template-assistant/{draft['id']}/apply", json=proposal, headers=headers)

    assert applied.status_code == 200
    schema = applied.json()["field_schema"]
    assert list(schema) == ["observaciones", "cliente", "fecha", "materiales"]  # its own fields first
    assert schema["cliente"]["label"] == "Cliente o comunidad"
    assert schema["materiales"]["columns"] == {"material": {"type": "str", "description": ""}, "cantidad": {"type": "int", "description": ""}}
    template = (await db.scalars(select(DocumentTemplate))).one()
    assert template.is_active and template.original_filename == "parte.docx"
    assert DocxTemplate(template.file_path).get_undeclared_template_variables() == {"cliente", "fecha", "materiales"}

    again = await client.post(f"/document-types/{doc_type.id}/template-assistant/{draft['id']}/apply", json=proposal, headers=headers)
    assert again.status_code == 404  # the draft is gone once applied


async def test_a_plan_that_cannot_work_is_refused_with_the_reason(client: AsyncClient, db: AsyncSession, model: AsyncMock) -> None:
    doc_type, headers = await _setup(client, db)
    draft = (await _upload(client, doc_type, headers, _example())).json()
    proposal = draft["proposal"]
    proposal["values"][0]["field"] = "materiales"  # a table field where one value goes

    refused = await client.post(f"/document-types/{doc_type.id}/template-assistant/{draft['id']}/apply", json=proposal, headers=headers)

    assert refused.status_code == 400 and "materiales" in refused.json()["detail"]
    assert (await db.scalars(select(DocumentTemplate))).all() == []


async def test_a_document_that_is_a_template_already_needs_no_model(
    client: AsyncClient, db: AsyncSession, model: AsyncMock
) -> None:
    doc_type, headers = await _setup(client, db)

    draft = (await _upload(client, doc_type, headers, _example(tagged=True))).json()

    model.assert_not_awaited()
    assert draft["already_tagged"] is True
    fields = {f["name"]: f for f in draft["proposal"]["fields"]}
    assert fields["cliente"]["type"] == "str"
    assert fields["materiales"]["type"] == "list[object]"
    assert [c["name"] for c in fields["materiales"]["columns"]] == ["cantidad", "material"]
    applied = await client.post(f"/document-types/{doc_type.id}/template-assistant/{draft['id']}/apply", json=draft["proposal"], headers=headers)
    assert applied.status_code == 200


async def test_without_an_ai_model_the_admin_is_told_what_to_set_up(
    client: AsyncClient, db: AsyncSession, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(settings, "ANTHROPIC_API_KEY", "")
    monkeypatch.setattr(settings, "EXTRACTION_PROVIDER", "anthropic")
    doc_type, headers = await _setup(client, db)

    response = await _upload(client, doc_type, headers, _example())

    assert response.status_code == 400 and "Settings → AI" in response.json()["detail"]


async def test_a_file_that_is_not_word_is_refused(client: AsyncClient, db: AsyncSession, model: AsyncMock) -> None:
    doc_type, headers = await _setup(client, db)

    response = await _upload(client, doc_type, headers, b"%PDF-1.4 not a docx")

    assert response.status_code == 400 and ".docx" in response.json()["detail"]

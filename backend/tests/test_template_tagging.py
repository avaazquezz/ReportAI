"""A filled-in example document becomes a template by editing the Word file itself: values become
tags, repeated table rows a row loop, a list a paragraph loop — and the rest stays as it was."""

from pathlib import Path

import pytest
from docx import Document
from docx.shared import Pt, RGBColor
from docxtpl import DocxTemplate

from app.services.rendering.docx_render import fill_template
from app.services.templates.docx_tagging import TagPlan, apply_plan, example_values, read_blocks


@pytest.fixture
def example(tmp_path: Path) -> Path:
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Reformas García S.L."
    title = document.add_paragraph()
    title.add_run("PARTE DE TRABAJO").bold = True
    p = document.add_paragraph()
    p.add_run("Cliente: ").bold = True
    split = p.add_run("Comunidad ")  # Word often splits one value across runs
    split.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)
    p.add_run("Calle Mayor 12").font.size = Pt(12)
    document.add_paragraph("Fecha: 14/09/2026    Técnico: Luis Pérez")
    document.add_paragraph("Trabajos realizados")
    document.add_paragraph("Cambio de bajante", style="List Bullet")
    document.add_paragraph("Sellado de juntas", style="List Bullet")
    table = document.add_table(rows=3, cols=3)
    for c, text in enumerate(["Material", "Cantidad", "Precio"]):
        table.rows[0].cells[c].text = text
    for r, row in enumerate([["Tubo PVC 110", "3", "12,50"], ["Silicona", "2", "6,00"]], start=1):
        for c, text in enumerate(row):
            table.rows[r].cells[c].text = text
    document.add_paragraph("Firma del cliente")
    path = tmp_path / "example.docx"
    document.save(str(path))
    return path


def _ids(path: Path) -> dict[str, str]:
    blocks, _ = read_blocks(str(path))
    return {b.text: b.id for b in blocks}


PLAN_FIELDS = {"cliente", "fecha", "tecnico", "trabajos", "materiales", "empresa"}


def _plan(path: Path) -> TagPlan:
    ids = _ids(path)
    return TagPlan(
        values=[
            (ids["Cliente: Comunidad Calle Mayor 12"], "Comunidad Calle Mayor 12", "cliente"),
            (ids["Fecha: 14/09/2026    Técnico: Luis Pérez"], "14/09/2026", "fecha"),
            (ids["Fecha: 14/09/2026    Técnico: Luis Pérez"], "Luis Pérez", "tecnico"),
            (ids["Reformas García S.L."], "Reformas García S.L.", "empresa"),
        ],
        lists=[([ids["Cambio de bajante"], ids["Sellado de juntas"]], "trabajos")],
        tables=[("t0", 1, 2, {0: "material", 1: "cantidad", 2: "precio"}, "materiales")],
    )


def test_the_document_is_read_as_addressable_blocks_and_tables(example: Path) -> None:
    blocks, tables = read_blocks(str(example))

    texts = {b.id: b.text for b in blocks}
    assert texts["h0p0"] == "Reformas García S.L." and texts["b1"] == "Cliente: Comunidad Calle Mayor 12"
    assert texts["t0r1c0p0"] == "Tubo PVC 110"
    assert tables[0].rows == [["Material", "Cantidad", "Precio"], ["Tubo PVC 110", "3", "12,50"], ["Silicona", "2", "6,00"]]


def test_tags_replace_the_values_and_the_rest_is_kept(example: Path, tmp_path: Path) -> None:
    out = tmp_path / "template.docx"

    problems = apply_plan(str(example), _plan(example), str(out))

    assert problems == []
    assert DocxTemplate(str(out)).get_undeclared_template_variables() == PLAN_FIELDS
    tagged = Document(str(out))
    texts = [p.text for p in tagged.paragraphs]
    assert "Cliente: {{ cliente }}" in texts
    assert "Fecha: {{ fecha }}    Técnico: {{ tecnico }}" in texts
    assert texts.count("{{ item }}") == 1 and "{%p for item in trabajos %}" in texts
    cliente = next(p for p in tagged.paragraphs if p.text.startswith("Cliente"))
    assert cliente.runs[0].bold and cliente.runs[1].font.color.rgb == RGBColor(0x1F, 0x2A, 0x44)
    assert "PARTE DE TRABAJO" in texts and "Firma del cliente" in texts
    assert tagged.sections[0].header.paragraphs[0].text == "{{ empresa }}"


def test_rendering_the_template_with_the_example_values_gives_back_the_example(example: Path, tmp_path: Path) -> None:
    plan = _plan(example)
    template = tmp_path / "template.docx"
    apply_plan(str(example), plan, str(template))

    rendered = Document(fill_template(str(template), example_values(str(example), plan), str(tmp_path / "out.docx")))

    texts = [p.text for p in rendered.paragraphs]
    assert "Cliente: Comunidad Calle Mayor 12" in texts
    assert "Fecha: 14/09/2026    Técnico: Luis Pérez" in texts
    assert [t for t in texts if t in ("Cambio de bajante", "Sellado de juntas")] == ["Cambio de bajante", "Sellado de juntas"]
    rows = [[cell.text for cell in row.cells] for row in rendered.tables[0].rows]
    assert rows == [["Material", "Cantidad", "Precio"], ["Tubo PVC 110", "3", "12,50"], ["Silicona", "2", "6,00"]]


def test_more_rows_than_the_example_are_rendered(example: Path, tmp_path: Path) -> None:
    template = tmp_path / "template.docx"
    apply_plan(str(example), _plan(example), str(template))
    materials = [{"material": f"M{i}", "cantidad": str(i), "precio": "1"} for i in range(5)]

    rendered = Document(fill_template(str(template), {"materiales": materials, "trabajos": []}, str(tmp_path / "out.docx")))

    assert len(rendered.tables[0].rows) == 6


def test_a_value_that_is_not_there_is_reported_not_dropped_silently(example: Path, tmp_path: Path) -> None:
    plan = TagPlan(values=[("b2", "Texto que no está", "cliente")])

    problems = apply_plan(str(example), plan, str(tmp_path / "out.docx"))

    assert problems == ["«Texto que no está» (cliente) was not found where it was"]

from pathlib import Path

import pytest
from docx import Document
from jinja2.exceptions import SecurityError
from PIL import Image
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.document_template import DocumentTemplate
from app.models.document_type import DocumentType
from app.models.report import Report
from app.models.report_attachment import ReportAttachment
from app.models.tenant import Tenant
from app.services.rendering.docx_render import fill_template
from app.services.rendering.report_document import (
    MissingTemplateError,
    assign_photos,
    fill_report_docx,
    template_values,
)


def _template(tmp_path: Path, text: str) -> str:
    path = tmp_path / "template.docx"
    doc = Document()
    doc.add_paragraph(text)
    doc.save(str(path))
    return str(path)


def _rendered_text(path: str) -> str:
    return "\n".join(p.text for p in Document(path).paragraphs)


def test_special_characters_in_values_survive(tmp_path: Path) -> None:
    template = _template(tmp_path, "Client: {{ name }}")
    out = fill_template(template, {"name": "García & Hijos <SL>"}, str(tmp_path / "out.docx"))
    assert _rendered_text(out) == "Client: García & Hijos <SL>"


def test_template_cannot_reach_python_internals(tmp_path: Path) -> None:
    # No template variables, so the upload-time validator lets it through; the sandbox must stop it.
    template = _template(tmp_path, "{{ ''.__class__.__mro__[1].__subclasses__() }}")
    with pytest.raises(SecurityError):
        fill_template(template, {}, str(tmp_path / "out.docx"))


def test_a_field_nobody_mentioned_prints_as_nothing(tmp_path: Path) -> None:
    template = _template(tmp_path, "Notes: {{ notes }}.")
    out = fill_template(template, {"notes": None}, str(tmp_path / "out.docx"))
    assert _rendered_text(out) == "Notes: ."


def test_a_list_nobody_mentioned_does_not_break_the_template(tmp_path: Path) -> None:
    schema = {"items": {"type": "list[str]"}, "rows": {"type": "list[object]", "columns": {"a": {"type": "str"}}}}
    values = template_values(schema, {"items": None, "rows": None}, "es")
    template = _template(tmp_path, "{% for i in items %}{{ i }}{% endfor %}{% for r in rows %}{{ r.a }}{% endfor %}end")
    assert _rendered_text(fill_template(template, values, str(tmp_path / "out.docx"))) == "end"


def test_photos_fill_the_slots_in_order_and_a_multiple_slot_takes_the_rest() -> None:
    schema = {
        "title": {"type": "str"},
        "cover": {"type": "image"},
        "gallery": {"type": "image", "multiple": True},
        "signature": {"type": "image"},
    }
    assert assign_photos(schema, ["a.jpg", "b.jpg", "c.jpg"]) == {
        "cover": "a.jpg", "gallery": ["b.jpg", "c.jpg"], "signature": None,
    }
    assert assign_photos(schema, []) == {"cover": None, "gallery": [], "signature": None}


def test_photos_are_placed_in_the_document(tmp_path: Path) -> None:
    photo = tmp_path / "photo.jpg"
    Image.new("RGB", (400, 300), "red").save(photo)
    template = _template(tmp_path, "{{ cover }}{% for p in gallery %}{{ p }}{% endfor %}{{ missing }}")

    out = fill_template(
        template, {}, str(tmp_path / "out.docx"),
        photos={"cover": str(photo), "gallery": [str(photo), str(photo)], "missing": None},
    )

    assert len(Document(out).inline_shapes) == 3


async def test_a_report_renders_with_its_photos_or_says_the_template_is_missing(
    db: AsyncSession, own_sessions: None, tmp_path: Path
) -> None:
    tenant = Tenant(name="Acme", slug="acme", is_active=True)
    db.add(tenant)
    await db.flush()
    schema = {"client": {"type": "str"}, "cover": {"type": "image"}}
    document_type = DocumentType(tenant_id=tenant.id, name="Visita", field_schema=schema)
    report = Report(tenant_id=tenant.id, status="pending", requester_channel="telegram", requester_identifier="1")
    db.add_all([document_type, report])
    await db.flush()
    photo = tmp_path / "photo.jpg"
    Image.new("RGB", (300, 400), "blue").save(photo)
    db.add(ReportAttachment(tenant_id=tenant.id, report_id=report.id, sender_identifier="1", path=str(photo)))
    await db.commit()
    render = {
        "report_id": report.id, "document_type_id": document_type.id, "field_schema": schema,
        "fields": {"client": "García"}, "language": "es", "output_path": str(tmp_path / "out.docx"),
    }

    with pytest.raises(MissingTemplateError):
        await fill_report_docx(**render)

    db.add(DocumentTemplate(
        tenant_id=tenant.id, document_type_id=document_type.id, version=1, is_active=True,
        file_path=_template(tmp_path, "{{ client }}{{ cover }}"),
    ))
    await db.commit()
    out = await fill_report_docx(**render)

    assert _rendered_text(out) == "García" and len(Document(out).inline_shapes) == 1


def test_dates_and_yes_no_read_as_a_person_writes_them() -> None:
    schema = {
        "day": {"type": "date"},
        "urgent": {"type": "bool"},
        "hours": {"type": "float"},
        "works": {"type": "list[object]", "columns": {"done": {"type": "date"}, "task": {"type": "str"}}},
    }
    fields = {"day": "2026-09-30", "urgent": False, "hours": 1.5, "works": [{"done": "2026-09-29", "task": "Caldera"}]}

    assert template_values(schema, fields, "es") == {
        "day": "30/09/2026", "urgent": "No", "hours": 1.5, "works": [{"done": "29/09/2026", "task": "Caldera"}],
    }
    assert template_values(schema, fields, "en")["urgent"] == "No"
    assert template_values(schema, fields, "en")["day"] == "2026-09-30"

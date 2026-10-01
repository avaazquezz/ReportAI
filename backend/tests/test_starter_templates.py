"""Ready-made document types: each builds a template whose tags are exactly its fields, in the
company's language, with its logo and name — and installs as a document type ready for reports."""

import uuid
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docxtpl import DocxTemplate
from httpx import AsyncClient
from PIL import Image
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.core.security import hash_password
from app.models.document_template import DocumentTemplate
from app.models.tenant import Tenant
from app.models.tenant_user import TenantUser
from app.schemas.document_type import DocumentTypeCreateRequest
from app.services.branding import Branding
from app.services.rendering.docx_render import fill_template
from app.services.rendering.report_document import assign_photos, template_values
from app.services.templates import starters


@pytest.mark.parametrize("language", ["es", "en"])
@pytest.mark.parametrize("key", list(starters.STARTERS))
def test_each_starter_renders_with_its_fields_logo_and_photos(key: str, language: str, tmp_path: Path) -> None:
    template = tmp_path / "template.docx"
    schema = starters.build(key, language, str(template))

    DocumentTypeCreateRequest(name="x", field_schema=schema)  # a valid schema for the panel
    tags = DocxTemplate(str(template)).get_undeclared_template_variables()
    assert tags == set(schema) | {"branding"}

    logo, photo = tmp_path / "logo.png", tmp_path / "photo.jpg"
    Image.new("RGB", (300, 100), "red").save(logo)
    Image.new("RGB", (400, 300), "blue").save(photo)
    sample: dict[str, object] = {}
    for name, spec in schema.items():
        if spec["type"] == "list[object]":
            sample[name] = [{column: "1" for column in spec["columns"]}]
        elif spec["type"].startswith("list["):
            sample[name] = ["uno", "dos"]
        elif spec["type"] == "enum":
            sample[name] = spec["options"][0]
        elif spec["type"] != "image":
            sample[name] = None if spec["type"] in ("date", "time", "bool") else "valor"
    out = fill_template(
        str(template), template_values(schema, sample, language), str(tmp_path / "out.docx"),
        photos=assign_photos(schema, [str(photo), str(photo)]),
        branding=Branding(name="Reformas García", color="#C0432A", logo_path=str(logo)),
    )

    rendered = Document(out)
    assert rendered.sections[0].header.paragraphs[1].text == "Reformas García"
    media = [n for n in zipfile.ZipFile(out).namelist() if n.startswith("word/media/")]
    assert len(media) >= 2  # the logo in the header, the photos in the body
    body = "\n".join(p.text for p in rendered.paragraphs)
    assert "{{" not in body and "{%" not in body


async def _admin(client: AsyncClient, db: AsyncSession) -> dict[str, str]:
    tenant = Tenant(name="Reformas García", slug=f"g-{uuid.uuid4().hex[:6]}", is_active=True, language="es")
    db.add(tenant)
    await db.flush()
    db.add(TenantUser(tenant_id=tenant.id, email="a@garcia.test", hashed_password=hash_password("pw-12345678"),
                      full_name="A", role="tenant_admin", is_active=True))
    await db.commit()
    login = await client.post("/auth/login", json={"email": "a@garcia.test", "password": "pw-12345678"})
    return {"Authorization": f"Bearer {login.json()['access_token']}"}


async def test_installing_a_starter_creates_a_ready_document_type(
    client: AsyncClient, db: AsyncSession, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(settings, "DOCUMENT_STORAGE_PATH", str(tmp_path))
    headers = await _admin(client, db)

    listed = (await client.get("/starter-templates", headers=headers)).json()
    assert [s["name"] for s in listed] == ["Parte de trabajo", "Informe de visita", "Informe de incidencia"]

    first = await client.post("/starter-templates/work_order/install", headers=headers)
    second = await client.post("/starter-templates/work_order/install", headers=headers)

    assert first.status_code == 201 and first.json()["name"] == "Parte de trabajo"
    assert second.json()["name"] == "Parte de trabajo (2)"
    assert first.json()["field_schema"]["trabajos"]["label"] == "Trabajos realizados"
    templates = (await db.scalars(select(DocumentTemplate))).all()
    assert len(templates) == 2 and all(t.is_active and Path(t.file_path).exists() for t in templates)
    assert (await client.post("/starter-templates/nope/install", headers=headers)).status_code == 404


def test_sections_with_nothing_in_them_are_left_out(tmp_path: Path) -> None:
    template = tmp_path / "template.docx"
    schema = starters.build("work_order", "es", str(template))

    out = fill_template(
        str(template),
        template_values(schema, {"cliente": "Acme", "trabajos": ["Revisión"]}, "es"),
        str(tmp_path / "out.docx"),
        photos=assign_photos(schema, []),
    )

    text = "\n".join(p.text for p in Document(out).paragraphs)
    assert "Trabajos realizados" in text
    assert "Fotos" not in text and "Observaciones" not in text and "Materiales" not in text
